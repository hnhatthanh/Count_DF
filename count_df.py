"""Count unique customer variables per FC (column CL) and status (Add/Change/Remove) in the 'Mapping' sheet."""
import ctypes
import sys
from copy import copy
from datetime import datetime
from pathlib import Path
from tkinter import Tk, filedialog

from openpyxl import load_workbook
from openpyxl.styles import Border, Side
import pandas as pd

STATUSES = ["Add", "Change", "Remove"]
COL_STATUS, COL_VAR, COL_FC = 0, 1, 89  # A, B, CL
COL_MAPPING_TYPE, COL_FC_NAME = 30, 39  # AE, AN
SW_RESTORE = 9


def count(path):
    df = pd.read_excel(path, sheet_name="Mapping", header=None, skiprows=3, usecols=[COL_STATUS, COL_VAR, COL_FC])
    df.columns = ["status", "var", "fc"]
    df["status"] = df["status"].astype(str).str.strip()
    df = df[df["status"] != "No change"]
    df = df.dropna(subset=["var"])
    # Excel's "Remove Duplicates" is case-insensitive
    df["key"] = df["var"].astype(str).str.strip().str.lower()
    df["fc"] = df["fc"].fillna("(Blanks)")
    report = (
        df[df["status"].isin(STATUSES)]
        .groupby(["fc", "status"])["key"].nunique()
        .unstack(fill_value=0)
        .reindex(columns=STATUSES, fill_value=0)
    )
    # Keep FCs that only have other non-"No change" statuses, like the Excel filter would
    report = report.reindex(sorted(df["fc"].unique(), key=str), fill_value=0)
    report.index.name = "FC Name"
    report.loc["Total"] = report.sum()
    return report


def count_mapping_fcs(data):
    data = data.copy()
    for column in ("status", "mapping_type", "fc"):
        data[column] = data[column].astype("string").str.strip()
    data["status"] = data["status"].str.casefold()
    data["mapping_type"] = data["mapping_type"].str.casefold()
    data = data[
        data["status"].ne("no change")
        & data["mapping_type"].isin(["complex", "factor/offset", "simple"])
        & data["fc"].notna()
        & data["fc"].ne("")
        & data["fc"].str.casefold().ne("none")
    ].copy()
    data["key"] = data["fc"].str.casefold()
    complex_fcs = data[data["mapping_type"].eq("complex")].drop_duplicates("key")["fc"].tolist()
    simple_fcs = data[data["mapping_type"].isin(["factor/offset", "simple"])].drop_duplicates("key")["fc"].tolist()
    return complex_fcs, simple_fcs


def write_report(report, path, bold_last=True):
    report.to_excel(path, sheet_name="Result", index=False)
    workbook = load_workbook(path)
    sheet = workbook["Result"]
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            if cell.row == 1 or (bold_last and cell.row == sheet.max_row):
                font = copy(cell.font)
                font.bold = True
                cell.font = font
    workbook.save(path)


def select_input_excel():
    root = Tk()
    try:
        root.withdraw()
        root.attributes("-topmost", True)
        root.update()
        return filedialog.askopenfilename(
            parent=root,
            title="Select input Excel",
            filetypes=[("Excel", "*.xls *.xlsx *.xlsm")],
        )
    finally:
        root.destroy()


def focus_program_window():
    if sys.platform != "win32":
        return
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    get_console_window = kernel32.GetConsoleWindow
    get_console_window.restype = ctypes.c_void_p
    window = get_console_window()
    if not window:
        return
    user32 = ctypes.WinDLL("user32", use_last_error=True)
    user32.ShowWindow.argtypes = (ctypes.c_void_p, ctypes.c_int)
    user32.SetForegroundWindow.argtypes = (ctypes.c_void_p,)
    user32.ShowWindow(window, SW_RESTORE)
    user32.SetForegroundWindow(window)


def main():
    print("[1] Count number of interfaces")
    print("[2] Count number FC")
    while True:
        mode = input("Select option: ").strip()
        if mode in ("1", "2"):
            break
        print("Enter 1 or 2.")

    if len(sys.argv) > 1:
        src = sys.argv[1]
    else:
        src = select_input_excel()
        if src:
            focus_program_window()
    if not src:
        sys.exit("No file selected.")
    src = Path(src)
    name = src.stem.split("_", 3)[-1] if src.stem.startswith("swsapsa_dataflow_") else src.stem
    if mode == "1":
        report = count(src)
        out = src.with_name(f"{name}_COUNTED_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        write_report(report.reset_index(), out)
        print(report.to_string())
        print(f"\nReport written to {out}")
        return

    data = pd.read_excel(
        src, sheet_name="Mapping", header=None, skiprows=3,
        usecols=[COL_STATUS, COL_MAPPING_TYPE, COL_FC_NAME],
    )
    data.columns = ["status", "mapping_type", "fc"]
    complex_fcs, simple_fcs = count_mapping_fcs(data)
    print(f"Complex FC: {len(complex_fcs)}")
    print(f"Simple FC: {len(simple_fcs)}")
    if input("Export report? (y/n) ").strip().lower() != "y":
        return

    report = pd.DataFrame(
        [("Complex", fc) for fc in complex_fcs] + [("Simple", fc) for fc in simple_fcs],
        columns=["Category", "FC Name"],
    )
    out = src.with_name(f"{name}_FC_COUNTED_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    write_report(report, out, bold_last=False)
    print(f"Report written to {out}")


if __name__ == "__main__":
    main()
