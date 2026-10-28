import ctypes
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook
import pandas as pd

import count_df


class MappingFcCountsTest(unittest.TestCase):
    def test_file_picker_is_topmost_and_owned(self):
        root = Mock()
        with patch.object(count_df, "Tk", return_value=root), patch.object(
            count_df.filedialog, "askopenfilename", return_value="book.xls"
        ) as ask:
            self.assertEqual(count_df.select_input_excel(), "book.xls")
        root.attributes.assert_called_once_with("-topmost", True)
        root.update.assert_called_once()
        root.destroy.assert_called_once()
        self.assertIs(ask.call_args.kwargs["parent"], root)

    def test_console_window_is_restored_and_activated(self):
        kernel32, user32 = Mock(), Mock()
        kernel32.GetConsoleWindow.return_value = 1234
        api = SimpleNamespace(c_void_p=ctypes.c_void_p, c_int=ctypes.c_int, WinDLL=Mock(side_effect=[kernel32, user32]))
        with patch("count_df.sys.platform", "win32"), patch.object(count_df, "ctypes", api, create=True):
            count_df.focus_program_window()
        user32.ShowWindow.assert_called_once_with(1234, 9)
        user32.SetForegroundWindow.assert_called_once_with(1234)

    def test_counts_complex_and_combined_simple_fc_names(self):
        data = pd.DataFrame([
            ("Add", "complex", "Alpha"),
            ("Change", "complex", "alpha"),
            ("No change", "complex", "Ignored"),
            ("Add", "factor/offset", "Beta"),
            ("Remove", "simple", "beta"),
            ("Change", "simple", "Gamma"),
            ("Add", "simple", "None"),
            ("Add", "simple", None),
            ("Add", "simple", " "),
            ("No change", "factor/offset", "Ignored simple"),
        ], columns=["status", "mapping_type", "fc"])

        self.assertTrue(callable(getattr(count_df, "count_mapping_fcs", None)))
        self.assertEqual(count_df.count_mapping_fcs(data), (["Alpha"], ["Beta", "Gamma"]))

    def test_fc_report_leaves_last_row_unbolded(self):
        report = pd.DataFrame([("Complex", "Alpha"), ("Simple", "Beta")], columns=["Category", "FC Name"])
        with TemporaryDirectory() as folder:
            path = Path(folder) / "report.xlsx"
            count_df.write_report(report, path, bold_last=False)
            sheet = load_workbook(path)["Result"]
            self.assertTrue(sheet.cell(1, 1).font.bold)
            self.assertFalse(sheet.cell(sheet.max_row, 1).font.bold)

            count_df.write_report(report, path)
            sheet = load_workbook(path)["Result"]
            self.assertTrue(sheet.cell(sheet.max_row, 1).font.bold)


if __name__ == "__main__":
    unittest.main()
